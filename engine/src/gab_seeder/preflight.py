from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .archive import EnvironmentArchive

UNIVERSE_MAP = {
    "student": "Student",
    "applied ml and data scientist": "Applied_ML_and_data_scientist",
    "startup founder": "Startup_founder",
    "educator and instructional designer": "Educator_and_instructional_designer",
    "backend software engineer": "Backend_software_engineer",
    "legal and contracts analyst": "Legal_and_contracts_analyst",
    "luxury travel advisor": "Luxury_travel_advisor",
    "indie game designer": "Indie_game_designer",
    "life sciences researcher": "Life_Sciences_Researcher",
}
KNOWN_OVERRIDE = "6a84bc5ab47b41eee8a5a796"


def load_tasks(workbook_path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    tasks: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
        index = {str(value): number + 1 for number, value in enumerate(headers) if value is not None}
        first_row = next(
            (row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value is not None),
            None,
        )
        if first_row is None:
            continue
        universe = str(sheet.cell(first_row, index["universe"]).value)
        prompt = str(sheet.cell(first_row, index["prompt"]).value)
        environment = UNIVERSE_MAP.get(universe.casefold())
        warning = None
        if sheet.title == KNOWN_OVERRIDE or ("#1121" in prompt and "Gemma-3" in prompt):
            if environment == "Student":
                warning = "rubric says Student but prompt evidence maps to Applied ML"
            environment = "Applied_ML_and_data_scientist"
        criterion_column = index.get("Criterion ID")
        criterion_count = 0
        if criterion_column:
            criterion_count = sum(
                1
                for row in range(2, sheet.max_row + 1)
                if sheet.cell(row, criterion_column).value is not None
            )
        tasks.append(
            {
                "sheet_id": sheet.title,
                "universe": universe,
                "environment": environment,
                "cuj": sheet.cell(first_row, index["CUJ"]).value,
                "prompt": prompt,
                "criterion_count": criterion_count,
                "warning": warning,
            }
        )
    return tasks


def load_account_emails(workbook_path: str | Path) -> list[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    result = []
    for row in sheet.iter_rows(values_only=True):
        value = row[0] if row else None
        if isinstance(value, str) and "@" in value:
            result.append(value.strip())
    return result


def suggested_account_mapping(accounts: list[str]) -> dict[str, Any]:
    order = [
        "Student",
        "Applied_ML_and_data_scientist",
        "Startup_founder",
        "Educator_and_instructional_designer",
        "Backend_software_engineer",
        "Legal_and_contracts_analyst",
        "Luxury_travel_advisor",
        "Indie_game_designer",
        "Life_Sciences_Researcher",
    ]
    mapping = {
        persona: {"email": email, "timezone": None}
        for persona, email in zip(order, accounts, strict=False)
    }
    spare = accounts[len(order) :]
    return {"accounts": mapping, "spares": spare}


def run_preflight(
    archive_path: str | Path,
    rubrics_path: str | Path,
    accounts_path: str | Path,
) -> dict[str, Any]:
    archive = EnvironmentArchive(archive_path)
    tasks = load_tasks(rubrics_path)
    accounts = load_account_emails(accounts_path)
    relevant = sorted({task["environment"] for task in tasks if task["environment"]})
    scans = [archive.scan_persona(persona) for persona in relevant]
    totals = {
        "calendar_events": sum(scan["calendar_events"] for scan in scans),
        "email_messages": sum(scan["email_messages"] for scan in scans),
        "filesystem_files": sum(scan["filesystem_files"] for scan in scans),
        "filesystem_declared_bytes": sum(scan["filesystem_declared_bytes"] for scan in scans),
        "missing_attachment_names": sum(len(scan["missing_attachment_names"]) for scan in scans),
        "name_only_attendee_values": sum(scan["name_only_attendee_values"] for scan in scans),
        "events_with_recurrence_hint": sum(scan["events_with_recurrence_hint"] for scan in scans),
        "github_files": sum(scan["github_files"] for scan in scans),
        "github_declared_bytes": sum(scan["github_declared_bytes"] for scan in scans),
        "github_worktree_files": sum(scan["github_worktree_files"] for scan in scans),
        "github_worktree_bytes": sum(scan["github_worktree_bytes"] for scan in scans),
        "github_git_store_files": sum(scan["github_git_store_files"] for scan in scans),
        "github_git_store_bytes": sum(scan["github_git_store_bytes"] for scan in scans),
    }
    task_counts = Counter(task["environment"] for task in tasks)
    return {
        "archive": str(Path(archive_path).expanduser().resolve()),
        "archive_personas": archive.personas(),
        "task_count": len(tasks),
        "task_environment_counts": dict(task_counts),
        "tasks": tasks,
        "account_count": len(accounts),
        "account_emails": accounts,
        "suggested_mapping": suggested_account_mapping(accounts),
        "relevant_personas": relevant,
        "persona_scans": scans,
        "totals": totals,
    }


def write_report(path: str | Path, report: dict[str, Any]) -> None:
    destination = Path(path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
